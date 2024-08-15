import os
import json
import requests
from msal import ConfidentialClientApplication
from docx import Document
from io import BytesIO
from dotenv import load_dotenv

# Load environment variables
load_dotenv()

CLIENT_ID = os.getenv('CLIENT_ID')
CLIENT_SECRET = os.getenv('CLIENT_SECRET')
TENANT_ID = os.getenv('TENANT_ID')
DRIVE_ID = os.getenv('DRIVE_ID')
ITEM_ID = os.getenv('ITEM_ID')

# Initialize the MSAL confidential client
app = ConfidentialClientApplication(
    CLIENT_ID,
    authority=f"https://login.microsoftonline.com/{TENANT_ID}",
    client_credential=CLIENT_SECRET,
)

# Acquire an access token for the Graph API
result = app.acquire_token_for_client(scopes=["https://graph.microsoft.com/.default"])

if "access_token" in result:
    access_token = result["access_token"]
else:
    print(result.get("error"))
    print(result.get("error_description"))
    exit(1)

# Define the endpoint to get the document content
endpoint = f"https://graph.microsoft.com/v1.0/drives/{DRIVE_ID}/items/{ITEM_ID}/content"
headers = {
    'Authorization': 'Bearer ' + access_token
}

# Download the document content
response = requests.get(endpoint, headers=headers)
response.raise_for_status()
document = Document(BytesIO(response.content))

# Replace placeholders in the document
def replace_placeholders(doc, replacements):
    for paragraph in doc.paragraphs:
        for placeholder, replacement in replacements.items():
            if placeholder in paragraph.text:
                paragraph.text = paragraph.text.replace(placeholder, replacement)
    return doc

# Define replacements dynamically
def get_replacements():
    # Example: Read replacements from a JSON file or another source
    # with open('replacements.json', 'r') as file:
    #     replacements = json.load(file)
    
    # For demonstration, using hardcoded replacements here
    replacements = {
        '{placeholder1}': 'DynamicReplacement1',
        '{placeholder2}': 'DynamicReplacement2',
        '{placeholder3}': 'DynamicReplacement3'
    }
    return replacements

# Get dynamic replacements
replacements = get_replacements()

# Modify the document
modified_document = replace_placeholders(document, replacements)

# Save the modified document to a BytesIO object
output_stream = BytesIO()
modified_document.save(output_stream)
output_stream.seek(0)

# Define the endpoint to upload the modified document
upload_endpoint = f"https://graph.microsoft.com/v1.0/drives/{DRIVE_ID}/items/{ITEM_ID}/content"

# Upload the modified document
upload_headers = {
    'Authorization': 'Bearer ' + access_token,
    'Content-Type': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
}
upload_response = requests.put(upload_endpoint, headers=upload_headers, data=output_stream)
upload_response.raise_for_status()

print("Document modified and uploaded successfully.")
